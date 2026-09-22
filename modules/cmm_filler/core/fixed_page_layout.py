"""
固定 A4 分页版式
================
按模板中各 Sheet 的序号列（A 列等）定位行，将实测值写入对应分页；
不复制 Sheet、不扩展行数。
"""

from __future__ import annotations

import logging
import re
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    import openpyxl

logger = logging.getLogger('CMMFiller')

LAYOUT_DYNAMIC = 'dynamic'
LAYOUT_FIXED_PAGES = 'fixed_pages'


def expand_serial_aliases(key: str) -> list[str]:
    """34.1 ↔ 34-1 等等价序号别名。"""
    aliases = [key]
    if re.fullmatch(r'\d+\.\d+', key):
        base, sub = key.split('.', 1)
        aliases.append(f'{base}-{sub}')
    elif re.fullmatch(r'\d+-\d+', key):
        base, sub = key.split('-', 1)
        aliases.append(f'{base}.{sub}')
    return list(dict.fromkeys(aliases))


def serial_cell_keys(value) -> list[str]:
    """将模板单元格中的序号转为可匹配键列表。"""
    if value is None:
        return []
    keys: list[str] = []
    if isinstance(value, (int, float)):
        if float(value) == int(value):
            keys.append(str(int(value)))
        else:
            s = str(value)
            keys.append(s)
            keys.append(f'{value:g}')
    else:
        s = str(value).strip()
        if not s:
            return []
        keys.append(s)
        try:
            f = float(s)
            if f == int(f):
                keys.append(str(int(f)))
            else:
                keys.append(f'{f:g}')
        except ValueError:
            pass
    out: list[str] = []
    for k in keys:
        out.extend(expand_serial_aliases(k))
    return list(dict.fromkeys(out))


def measure_lookup_keys(measure: dict) -> list[str]:
    """从 PDF 解析结果生成模板序号查找键（优先 34.1 风格）。"""
    num = measure['num']
    sub = measure.get('sub') or ''
    sub_sep = measure.get('sub_sep', '-')
    keys: list[str] = []

    if not sub:
        keys.append(str(num))
    elif sub.isdigit():
        if sub_sep == '.':
            keys.append(f'{num}.{sub}')
        elif sub_sep == '_':
            keys.append(f'{num}_{sub}')
            keys.append(f'{num}.{sub}')
        else:
            keys.append(f'{num}.{sub}')
            keys.append(f'{num}-{sub}')
    else:
        keys.append(measure.get('item_key') or f'{num}{sub_sep}{sub}')
        keys.append(f'{num}.{sub}')

    if measure.get('item_key'):
        keys.append(measure['item_key'])

    out: list[str] = []
    for k in keys:
        out.extend(expand_serial_aliases(k))
    return list(dict.fromkeys(out))


def build_sheet_row_index(
    wb: openpyxl.Workbook,
    sheet_names: list[str],
    serial_col: int,
    data_start_row: int,
    max_data_row: int,
) -> dict[str, tuple[str, int]]:
    """
    扫描各 Sheet 序号列，建立 lookup_key -> (sheet_name, row) 索引。
    同一键在多个 Sheet 出现时，保留首次（通常不应重复）。
    """
    index: dict[str, tuple[str, int]] = {}
    for sheet_name in sheet_names:
        if sheet_name not in wb.sheetnames:
            logger.warning(f'[固定分页] 模板缺少 Sheet "{sheet_name}"，已跳过')
            continue
        ws = wb[sheet_name]
        for row in range(data_start_row, max_data_row + 1):
            value = ws.cell(row=row, column=serial_col).value
            for key in serial_cell_keys(value):
                if key not in index:
                    index[key] = (sheet_name, row)
    return index


def resolve_measure_location(
    row_index: dict[str, tuple[str, int]],
    measure: dict,
) -> tuple[str, int] | None:
    for key in measure_lookup_keys(measure):
        if key in row_index:
            return row_index[key]
    return None


def load_axis_preferences_from_template(
    template_path: str,
    sheet_names: list[str],
    serial_col: int,
    axis_col: int,
    data_start_row: int,
    max_data_row: int,
) -> dict[str, str]:
    """
    扫描模板各 Sheet 的「轴」列，建立 序号 → 轴字母 映射。
    用户在模板对应行填入 D / X / Y / M / A 等即可指定多轴项取哪一行。
    """
    import openpyxl

    prefs: dict[str, str] = {}
    try:
        wb = openpyxl.load_workbook(template_path, data_only=True, read_only=True)
    except Exception as e:
        logger.warning(f'[轴偏好] 无法读取模板: {e}')
        return prefs

    try:
        for sheet_name in sheet_names:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            for row in range(data_start_row, max_data_row + 1):
                raw = ws.cell(row=row, column=axis_col).value
                if raw is None:
                    continue
                axis = str(raw).strip().upper()
                if not (len(axis) == 1 and axis.isalpha()):
                    continue
                serial = ws.cell(row=row, column=serial_col).value
                for key in serial_cell_keys(serial):
                    prefs[key] = axis
    finally:
        wb.close()

    if prefs:
        logger.info(f'[轴偏好] 已从模板加载 {len(prefs)} 项轴选择')
    return prefs
