"""
通用 PDF 表格提取（可选依赖 pdfplumber）
========================================
仅用于文字型 PDF 的表格导入，扫描件请走 OCR 流程。
"""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

logger = logging.getLogger('CMMFiller')

_HEADER_FILL = PatternFill('solid', fgColor='D9D9D9')
_HEADER_FONT = Font(bold=True)


def pdfplumber_available() -> bool:
    try:
        import pdfplumber  # noqa: F401
        return True
    except ImportError:
        return False


def extract_tables_from_pdf(pdf_path: str | Path) -> list[dict]:
    """
    从文字型 PDF 提取所有表格。

    Returns:
        [{'page': int, 'table_idx': int, 'rows': list[list[str]]}, ...]
    """
    import pdfplumber

    tables_out: list[dict] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page_i, page in enumerate(pdf.pages):
            tables = page.extract_tables() or []
            for table_i, tbl in enumerate(tables):
                if not tbl:
                    continue
                cleaned = []
                for row in tbl:
                    cleaned.append([
                        (cell.strip() if isinstance(cell, str) else (str(cell) if cell is not None else ''))
                        for cell in row
                    ])
                tables_out.append({
                    'page': page_i + 1,
                    'table_idx': table_i,
                    'rows': cleaned,
                })
    return tables_out


def export_tables_to_excel(tables: list[dict], output_path: str | Path, source_name: str = '') -> str:
    """将提取的表格写入 Excel，每表一个 Sheet。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()

    for i, tbl in enumerate(tables):
        title = f'P{tbl["page"]}_T{tbl["table_idx"] + 1}'
        if source_name:
            base = Path(source_name).stem[:20]
            title = f'{base}_{title}'
        title = title[:31]
        n = 2
        orig = title
        while title.lower() in used_titles:
            suffix = f'_{n}'
            title = orig[:31 - len(suffix)] + suffix
            n += 1
        used_titles.add(title.lower())

        ws = wb.create_sheet(title)
        for r, row in enumerate(tbl['rows'], 1):
            for c, val in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=val)
                if r == 1:
                    cell.fill = _HEADER_FILL
                    cell.font = _HEADER_FONT
                    cell.alignment = Alignment(horizontal='center')

    if not wb.sheetnames:
        ws = wb.create_sheet('空')
        ws.cell(row=1, column=1, value='未提取到表格')

    wb.save(str(output_path))
    return str(output_path)


def batch_export_pdf_tables(pdf_paths: list[str | Path], output_folder: str | Path) -> dict:
    """
    批量提取多个 PDF 的表格，每个 PDF 输出一个 Excel。

    Returns summary dict.
    """
    if not pdfplumber_available():
        raise ImportError(
            'pdfplumber 未安装。请运行: pip install pdfplumber\n'
            '（仅适用于文字型 PDF，扫描件请使用 OCR 流程）'
        )

    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)
    summary = {'total': len(pdf_paths), 'processed': 0, 'failed': [], 'output_files': []}

    for pdf_path in pdf_paths:
        pdf_path = Path(pdf_path)
        try:
            tables = extract_tables_from_pdf(pdf_path)
            out = output_folder / f'{pdf_path.stem}_表格.xlsx'
            export_tables_to_excel(tables, out, source_name=pdf_path.name)
            summary['processed'] += 1
            summary['output_files'].append(str(out))
            logger.info(f'[表格提取] {pdf_path.name}: {len(tables)} 个表格 → {out.name}')
        except Exception as e:
            logger.error(f'[表格提取失败] {pdf_path.name}: {e}')
            summary['failed'].append({'file': pdf_path.name, 'error': str(e)})

    return summary
