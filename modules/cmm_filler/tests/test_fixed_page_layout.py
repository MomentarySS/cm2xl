"""固定 A4 分页版式单元测试"""

import unittest
from pathlib import Path
from unittest.mock import MagicMock

from openpyxl import Workbook

from modules.cmm_filler.core.fixed_page_layout import (
    build_sheet_row_index,
    load_axis_preferences_from_template,
    measure_lookup_keys,
    resolve_measure_location,
    serial_cell_keys,
)
from modules.cmm_filler.core.filler import CMMReportFiller


class TestFixedPageLayout(unittest.TestCase):
    def test_serial_cell_keys_decimal_sub(self):
        self.assertIn('34.1', serial_cell_keys(34.1))
        self.assertIn('34-1', serial_cell_keys(34.1))

    def test_measure_lookup_keys_hyphen_sub(self):
        keys = measure_lookup_keys({
            'num': 34, 'sub': '1', 'sub_sep': '-', 'item_key': '34-1',
        })
        self.assertIn('34.1', keys)
        self.assertIn('34-1', keys)

    def test_build_sheet_row_index(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'A'
        ws['A6'] = 1
        ws['A7'] = 34.1
        ws['A8'] = 34.2
        index = build_sheet_row_index(wb, ['A'], serial_col=1, data_start_row=6, max_data_row=10)
        self.assertEqual(index['1'], ('A', 6))
        self.assertEqual(index['34.1'], ('A', 7))
        self.assertEqual(resolve_measure_location(index, {
            'num': 34, 'sub': '2', 'sub_sep': '-', 'item_key': '34-2',
        }), ('A', 8))

    def test_main_sample_count_from_config(self):
        filler = MagicMock(spec=CMMReportFiller)
        filler.config = {'main_sample_count': 3}
        self.assertEqual(CMMReportFiller._get_main_sample_count(filler), 3)

    def test_batch_output_filename(self):
        self.assertEqual(
            CMMReportFiller._batch_output_filename('2026-08-28', 1, 6, multi_batch=False),
            '2026-08-28_项目汇总.xlsx',
        )
        self.assertEqual(
            CMMReportFiller._batch_output_filename('2026-08-28', 1, 6, multi_batch=True),
            '2026-08-28_项目汇总_样品1-6.xlsx',
        )
        self.assertEqual(
            CMMReportFiller._batch_output_filename('2026-08-28', 7, 10, multi_batch=True),
            '2026-08-28_项目汇总_样品7-10.xlsx',
        )

    def test_prepare_sample_batches_splits_when_over_capacity(self):
        filler = MagicMock(spec=CMMReportFiller)
        filler._overflow_samples_enabled = lambda: False
        filler._chunk_samples = CMMReportFiller._chunk_samples.__get__(filler, CMMReportFiller)
        items = [{'path': i} for i in range(10)]
        batches = CMMReportFiller._prepare_sample_batches(filler, items, 6)
        self.assertEqual(len(batches), 2)
        self.assertEqual(len(batches[0]), 6)
        self.assertEqual(len(batches[1]), 4)

    def test_renumber_batch_items(self):
        batch = [{'stem': 'a', 'sample_num': 7}, {'stem': 'b', 'sample_num': 8}]
        out = CMMReportFiller._renumber_batch_items(batch)
        self.assertEqual(out[0]['sample_num'], 1)
        self.assertEqual(out[1]['sample_num'], 2)
        self.assertEqual(out[0]['stem'], 'a')

    def test_load_axis_preferences_from_workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'A'
        ws['A6'] = 15
        ws['G6'] = 'X'
        ws['A7'] = 16
        ws['G7'] = 'D'
        path = self._tmpdir / 'axis_tpl.xlsx'
        wb.save(path)
        wb.close()
        prefs = load_axis_preferences_from_template(
            str(path), ['A'], serial_col=1, axis_col=7,
            data_start_row=6, max_data_row=10,
        )
        self.assertEqual(prefs.get('15'), 'X')
        self.assertEqual(prefs.get('16'), 'D')

    @classmethod
    def setUpClass(cls):
        import tempfile
        cls._tmpdir = Path(tempfile.mkdtemp())


if __name__ == '__main__':
    unittest.main()
