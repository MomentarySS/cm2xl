# -*- coding: utf-8 -*-
"""汇总导出回归测试：多 PDF → 单 Excel，每个 PDF 一个 Sheet。

核心层断言：sheet 顺序、标题/信息行/表头、逐列数据与 OCR 解析结果一致、
NG 行红字与趋势列计算、数字格式、冻结表头、sheet 名净化。
最后附 GUI 冒烟（实例化 + 文件列表增删，不进 mainloop）。
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from cmm_filler_v10 import (
    CMMReportFiller, TEMPLATE_PATH, SUMMARY_HEADERS, _safe_sheet_title,
)

SAMPLES = Path(__file__).resolve().parent / 'samples'
OUTPUT = Path(__file__).resolve().parent / 'outputs' / 'test_summary'


def _is_red(cell):
    """openpyxl 颜色可能是 FF0000 / 00FF0000 / FFFF0000，统一按末 6 位判断"""
    color = cell.font.color
    return bool(color and getattr(color, 'rgb', None) and str(color.rgb).upper().endswith('FF0000'))


def main():
    filler = CMMReportFiller(TEMPLATE_PATH, dpi=300)

    all_pdfs = sorted(SAMPLES.glob('*.PDF'))
    assert all_pdfs, 'samples 目录为空'

    # 优先用已缓存（OCR 过）的 PDF，保证测试快
    cached_stems = {k[:-4] for k in filler.cache.keys()}  # '0624-001.png' -> stem
    cached_pdfs = [p for p in all_pdfs if p.stem in cached_stems]

    # 从缓存里找含 NG 项的 PDF（用于验证红字/趋势列）
    ng_pdfs = []
    for p in cached_pdfs:
        lines = filler.cache.get(p.stem + '.png')
        if not lines:
            continue
        data = filler.parse_from_text(lines)
        if any(m.get('ng') for m in data.get('measurements', [])):
            ng_pdfs.append(p)
    print(f'缓存 PDF: {len(cached_pdfs)}/{len(all_pdfs)}，其中含 NG 项的: {len(ng_pdfs)}')

    # 选 8 个：不同日期组各挑 1 个（共 6 个）+ 含 NG 的最多 2 个（去重）
    picked = []
    seen_dates = set()
    for p in cached_pdfs:
        lines = filler.cache.get(p.stem + '.png')
        if not lines:
            continue
        data = filler.parse_from_text(lines)
        d = data.get('date', '')
        if d and d not in seen_dates:
            seen_dates.add(d)
            picked.append(p)
        if len(picked) >= 6:
            break
    for p in ng_pdfs[:2]:
        if p not in picked:
            picked.append(p)
    picked = picked[:8]
    assert len(picked) >= 6, f'可测样本不足: {len(picked)}'
    print('选中:', [p.name for p in picked])

    summary = filler.export_summary_workbook([str(p) for p in picked], str(OUTPUT))
    print('summary:', {k: v for k, v in summary.items() if k != 'failed_pdfs'})
    assert summary['total_pdfs'] == len(picked)
    assert summary['processed_pdfs'] == len(picked)
    assert not summary['failed_pdfs'], summary['failed_pdfs']

    import openpyxl
    wb = openpyxl.load_workbook(summary['output_file'])
    assert wb.sheetnames == [p.stem for p in picked], wb.sheetnames
    print('sheet 顺序 OK:', wb.sheetnames)

    ng_checked = 0
    for p in picked:
        ws = wb[p.stem]
        data = filler.parse_from_text(filler.cache[p.stem + '.png'])
        ms = sorted(data['measurements'], key=lambda m: m['num'])

        # 标题 / 信息行 / 表头 / 冻结
        assert ws['A1'].value == f'CMM Dimensions Report 汇总 — {p.name}', ws['A1'].value
        assert ws['A2'].value.startswith('零件名: ') and f'序列号: {p.stem}' in ws['A2'].value, ws['A2'].value
        headers = [ws.cell(row=3, column=c).value for c in range(1, 11)]
        assert headers == SUMMARY_HEADERS, headers
        assert 'A1:J1' in {str(r) for r in ws.merged_cells.ranges}
        assert ws.freeze_panes == 'A4'

        assert ws.max_row == 3 + len(ms), (p.name, ws.max_row, len(ms))
        for i, m in enumerate(ms):
            r = 4 + i
            dev = round(m['measured'] - m['nominal'], 4)
            vals = [ws.cell(row=r, column=c).value for c in range(1, 11)]
            assert vals[0] == m['num'], (p.name, r, vals[0])
            assert vals[1] == m['desc'], (p.name, r, vals[1])
            assert vals[2] is None and vals[3] is None  # 坐标系/补偿留空
            assert abs(vals[4] - m['nominal']) < 1e-9
            assert abs(vals[5] - m['measured']) < 1e-9
            assert abs(vals[6] - dev) < 1e-9, (p.name, r, vals[6], dev)
            assert abs(vals[7] - m['lower_tol']) < 1e-9
            assert abs(vals[8] - m['upper_tol']) < 1e-9
            assert ws.cell(row=r, column=6).number_format == '0.0000'
            if m.get('ng'):
                expect = round(dev - m['upper_tol'], 4) if dev > m['upper_tol'] else round(dev - m['lower_tol'], 4)
                assert vals[9] is not None and abs(vals[9] - expect) < 1e-9, (p.name, r, vals[9], expect)
                for col in (6, 7, 10):
                    assert _is_red(ws.cell(row=r, column=col)), (p.name, r, col)
                ng_checked += 1
            else:
                assert vals[9] is None, (p.name, r, vals[9])
                assert not _is_red(ws.cell(row=r, column=6)), (p.name, r)
    print(f'逐 sheet 数据核对通过；NG 红字行验证 {ng_checked} 处（总 NG 计 {summary["ng_count"]}）')
    assert ng_checked == summary['ng_count'], (ng_checked, summary['ng_count'])

    # sheet 名净化：重名后缀 / 非法字符 / 31 字符截断
    used = set()
    assert _safe_sheet_title('0624-001', used) == '0624-001'
    assert _safe_sheet_title('0624-001', used) == '0624-001_2'
    assert _safe_sheet_title('a[b]:*?/\\c', set()) == 'a_b______c'
    assert len(_safe_sheet_title('x' * 50, set())) == 31

    print('\n=== 核心层测试全部通过 ===')

    # GUI 冒烟：实例化 + 文件列表管理（不进 mainloop、不真正导出）
    from cmm_filler_gui import CMMFillerGUI
    app = CMMFillerGUI()
    app._summary_extend_files([str(p) for p in picked[:3]])
    assert len(app._summary_files) == 3, len(app._summary_files)
    app._summary_extend_files([str(picked[0])])  # 重复添加应被忽略
    assert len(app._summary_files) == 3
    app._summary_extend_files([str(p) for p in picked[3:5]])
    assert len(app._summary_files) == 5
    app._summary_remove_file(0)
    assert len(app._summary_files) == 4
    assert app.summary_count_label.cget('text') == '已选 4 个 PDF'
    app._summary_clear_files()
    assert len(app._summary_files) == 0
    app.root.destroy()
    print('=== GUI 冒烟测试通过 ===')


if __name__ == '__main__':
    main()
