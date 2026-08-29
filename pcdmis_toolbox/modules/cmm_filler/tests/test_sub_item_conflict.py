"""子编号冲突检测与解析测试"""

import unittest

from openpyxl import Workbook

from modules.cmm_filler.core.fixed_page_layout import build_sheet_row_index
from modules.cmm_filler.core.sub_item_conflict import (
    SKIP_SUB_ITEM,
    WORST_NG_SUB_ITEM,
    build_parent_auto_picks,
    detect_parent_row_conflicts,
    pick_worst_ng_candidate,
    resolve_measure_write_location,
    resolve_parent_picks,
)


def _row_index_parent_only() -> dict[str, tuple[str, int]]:
    wb = Workbook()
    ws = wb.active
    ws.title = 'A'
    ws['A6'] = 1
    ws['A7'] = 2
    return build_sheet_row_index(wb, ['A'], serial_col=1, data_start_row=6, max_data_row=10)


def _m(num, sub, measured, ng=False, **kw) -> dict:
    item_key = f'{num}-{sub}' if sub else str(num)
    return {
        'num': num,
        'sub': sub,
        'sub_sep': '-',
        'item_key': item_key,
        'label': f'FAI_{num}-{sub}' if sub else f'FAI_{num:02d}',
        'nominal': 10.0,
        'upper_tol': 0.1,
        'lower_tol': -0.1,
        'measured': measured,
        'ng': ng,
        **kw,
    }


class TestSubItemConflict(unittest.TestCase):
    def test_detect_many_sub_items(self):
        row_index = _row_index_parent_only()
        measurements = [_m(1, str(i), 10.0 + i * 0.01) for i in range(1, 11)]
        conflicts = detect_parent_row_conflicts(measurements, row_index)
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(conflicts[0]['template_key'], '1')
        self.assertEqual(len(conflicts[0]['candidates']), 10)

    def test_no_conflict_when_single_sub(self):
        row_index = _row_index_parent_only()
        measurements = [_m(1, '1', 10.01)]
        self.assertEqual(detect_parent_row_conflicts(measurements, row_index), [])
        auto = build_parent_auto_picks(measurements, row_index)
        self.assertEqual(auto, {'1': '1-1'})

    def test_pick_worst_ng_candidate(self):
        candidates = [
            _m(1, '1', 10.01, ng=False),
            _m(1, '2', 10.5, ng=True),
            _m(1, '3', 10.02, ng=False),
        ]
        worst = pick_worst_ng_candidate(candidates)
        self.assertEqual(worst['sub'], '2')

    def test_resolve_parent_picks_skip_and_worst(self):
        conflict = {
            'template_key': '1',
            'candidates': [_m(1, '1', 10.01), _m(1, '2', 10.5, ng=True)],
        }
        skip = resolve_parent_picks([conflict], {'1': SKIP_SUB_ITEM})
        self.assertIsNone(skip['1'])
        worst = resolve_parent_picks([conflict], {'1': WORST_NG_SUB_ITEM})
        self.assertEqual(worst['1'], '1-2')
        explicit = resolve_parent_picks([conflict], {'1': '1-1'})
        self.assertEqual(explicit['1'], '1-1')

    def test_resolve_measure_write_location(self):
        row_index = _row_index_parent_only()
        parent_picks = {'1': '1-3', '2': None}
        loc = resolve_measure_write_location(_m(1, '3', 10.03), row_index, parent_picks)
        self.assertEqual(loc, ('A', 6))
        self.assertIsNone(
            resolve_measure_write_location(_m(1, '1', 10.01), row_index, parent_picks)
        )
        self.assertIsNone(
            resolve_measure_write_location(_m(2, '1', 10.01), row_index, parent_picks)
        )


if __name__ == '__main__':
    unittest.main()
