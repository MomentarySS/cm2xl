"""NG 统计分析单元测试"""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock

from modules.cmm_filler.core.ng_analysis import (
    collect_ng_stats,
    export_ng_workbook,
    format_ng_summary_text,
    _severity,
)
from modules.cmm_filler.core.filler import CMMReportFiller


def _make_info(name, measurements):
    path = MagicMock()
    path.name = name
    return {'path': path, 'stem': Path(name).stem, 'data': {'measurements': measurements}}


class TestNGAnalysis(unittest.TestCase):
    def test_collect_ng_stats(self):
        infos = [
            _make_info('001.PDF', [
                {'num': 1, 'sub': '', 'item_key': '1', 'label': 'FAI_01', 'desc': '孔径',
                 'nominal': 10.0, 'upper_tol': 0.1, 'lower_tol': -0.1, 'measured': 10.5, 'ng': True},
                {'num': 2, 'sub': '', 'item_key': '2', 'label': 'FAI_02', 'desc': '长度',
                 'nominal': 5.0, 'upper_tol': 0.05, 'lower_tol': -0.05, 'measured': 5.01, 'ng': False},
            ]),
            _make_info('002.PDF', [
                {'num': 1, 'sub': '', 'item_key': '1', 'label': 'FAI_01', 'desc': '孔径',
                 'nominal': 10.0, 'upper_tol': 0.1, 'lower_tol': -0.1, 'measured': 10.3, 'ng': True},
            ]),
        ]
        stats = collect_ng_stats(infos)
        self.assertEqual(stats['total_reports'], 2)
        self.assertEqual(stats['total_ng'], 2)
        self.assertEqual(stats['fai_count'], 2)
        fai1 = next(r for r in stats['summary'] if r['item_key'] == '1')
        self.assertEqual(fai1['ng_count'], 2)
        self.assertAlmostEqual(fai1['ng_rate'], 1.0)
        self.assertEqual(len(stats['details']), 2)

    def test_severity_levels(self):
        self.assertEqual(_severity(0.35), '严重')
        self.assertEqual(_severity(0.15), '警告')
        self.assertEqual(_severity(0.05), '轻微')
        self.assertEqual(_severity(0), '正常')

    def test_export_workbook(self):
        stats = collect_ng_stats([
            _make_info('a.PDF', [
                {'num': 1, 'sub': 0, 'item_key': '1', 'label': 'FAI_01', 'desc': 'X',
                 'nominal': 1.0, 'upper_tol': 0.1, 'lower_tol': -0.1, 'measured': 2.0, 'ng': True},
            ]),
        ])
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / 'ng_test.xlsx'
            export_ng_workbook(stats, out)
            self.assertTrue(out.exists())
            import openpyxl
            wb = openpyxl.load_workbook(out)
            self.assertIn('NG统计汇总', wb.sheetnames)
            self.assertIn('NG详细记录', wb.sheetnames)
            self.assertIn('FAI明细', wb.sheetnames)

    def test_format_summary_text(self):
        stats = {'total_reports': 3, 'fai_count': 5, 'total_ng': 2, 'summary': [
            {'num': 1, 'sub': '', 'item_key': '1', 'label': 'FAI_01', 'desc': 'test',
             'ng_count': 2, 'total': 3, 'ng_rate': 0.67, 'severity': '严重'},
        ]}
        text = format_ng_summary_text(stats)
        self.assertIn('分析报告数: 3', text)
        self.assertIn('FAI_01', text)


class TestSampleChunking(unittest.TestCase):
    def test_chunk_samples(self):
        filler = CMMReportFiller.__new__(CMMReportFiller)
        items = list(range(12))
        chunks = filler._chunk_samples(items, 5)
        self.assertEqual(len(chunks), 3)
        self.assertEqual(len(chunks[0]), 5)
        self.assertEqual(len(chunks[1]), 5)
        self.assertEqual(len(chunks[2]), 2)


if __name__ == '__main__':
    unittest.main()
