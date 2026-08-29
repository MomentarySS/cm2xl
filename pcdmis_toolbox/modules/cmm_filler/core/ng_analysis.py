"""
CMM 报告 NG 统计分析
====================
从批量 PDF 解析结果聚合 NG 率排行，导出多 Sheet Excel 报告。
"""

from __future__ import annotations

import logging
import time
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .parse_measurements import measure_dict_key, measure_dict_label, _measurement_sort_key

logger = logging.getLogger('CMMFiller')

_HEADER_FILL = PatternFill('solid', fgColor='4472C4')
_HEADER_FONT = Font(bold=True, color='FFFFFF')
_SEVERE_FILL = PatternFill('solid', fgColor='FFC7CE')   # NG率 >= 30%
_WARN_FILL = PatternFill('solid', fgColor='FCE4D6')      # 10% - 30%
_MILD_FILL = PatternFill('solid', fgColor='FFEB9C')     # < 10% 但有 NG
_DETAIL_FILL = PatternFill('solid', fgColor='FFF2F2')
_CENTER = Alignment(horizontal='center', vertical='center')
_LEFT = Alignment(horizontal='left', vertical='center')


def _severity(ng_rate: float) -> str:
    if ng_rate >= 0.3:
        return '严重'
    if ng_rate >= 0.1:
        return '警告'
    if ng_rate > 0:
        return '轻微'
    return '正常'


def _severity_fill(ng_rate: float) -> PatternFill | None:
    if ng_rate >= 0.3:
        return _SEVERE_FILL
    if ng_rate >= 0.1:
        return _WARN_FILL
    if ng_rate > 0:
        return _MILD_FILL
    return None


def collect_ng_stats(infos: list[dict]) -> dict:
    """
    从 _scan_and_parse 的 infos 列表聚合 NG 统计。

    每项 info: {path, stem, data: {measurements: [...]}}
  """
    fai_stats: dict[int, dict] = {}
    details: list[dict] = []
    total_reports = 0
    total_ng = 0

    for info in infos:
        data = info.get('data', {})
        measurements = data.get('measurements', [])
        if not measurements:
            continue
        total_reports += 1
        pdf_name = info['path'].name if hasattr(info.get('path'), 'name') else str(info.get('path', ''))

        for m in measurements:
            ikey = measure_dict_key(m)
            label = measure_dict_label(m)
            is_ng = bool(m.get('ng'))
            if is_ng:
                total_ng += 1

            if ikey not in fai_stats:
                fai_stats[ikey] = {
                    'item_key': ikey,
                    'num': m['num'],
                    'sub': m.get('sub', ''),
                    'label': label,
                    'desc': m.get('desc', ''),
                    'nominal': m['nominal'],
                    'upper_tol': m['upper_tol'],
                    'lower_tol': m['lower_tol'],
                    'ng_count': 0,
                    'pass_count': 0,
                }
            stat = fai_stats[ikey]
            if is_ng:
                stat['ng_count'] += 1
            else:
                stat['pass_count'] += 1
            if len(m.get('desc', '')) > len(stat['desc']):
                stat['desc'] = m['desc']

            if is_ng:
                dev = round(m['measured'] - m['nominal'], 4)
                details.append({
                    'file': pdf_name,
                    'item_key': ikey,
                    'label': label,
                    'num': m['num'],
                    'sub': m.get('sub', ''),
                    'desc': m.get('desc', ''),
                    'measured': m['measured'],
                    'nominal': m['nominal'],
                    'upper_tol': m['upper_tol'],
                    'lower_tol': m['lower_tol'],
                    'deviation': dev,
                })

    summary_rows = []
    for ikey in sorted(fai_stats, key=lambda k: (fai_stats[k]['num'], fai_stats[k]['sub'])):
        s = fai_stats[ikey]
        total = s['ng_count'] + s['pass_count']
        ng_rate = s['ng_count'] / total if total else 0.0
        summary_rows.append({
            **s,
            'total': total,
            'ng_rate': ng_rate,
            'severity': _severity(ng_rate),
        })

    summary_rows.sort(key=lambda r: (-r['ng_rate'], -r['ng_count'], r['num'], r.get('sub', '')))

    return {
        'total_reports': total_reports,
        'total_ng': total_ng,
        'summary': summary_rows,
        'details': details,
        'fai_count': len(fai_stats),
    }


def export_ng_workbook(stats: dict, output_path: str | Path) -> str:
    """导出 NG 统计 Excel（3 个 Sheet）。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── Sheet 1: NG 统计汇总 ──
    ws1 = wb.create_sheet('NG统计汇总')
    headers1 = ['排名', 'FAI', '描述', 'NG次数', 'PASS次数', 'NG率', '严重程度', '名义值', '上公差', '下公差']
    for col, h in enumerate(headers1, 1):
        c = ws1.cell(row=1, column=col, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = _CENTER

    for rank, row in enumerate(stats['summary'], 1):
        r = rank + 1
        values = [
            rank,
            row.get('label', f'FAI_{row["num"]:02d}'),
            row['desc'],
            row['ng_count'],
            row['pass_count'],
            f'{row["ng_rate"]:.1%}',
            row['severity'],
            row['nominal'],
            row['upper_tol'],
            row['lower_tol'],
        ]
        fill = _severity_fill(row['ng_rate'])
        for col, val in enumerate(values, 1):
            c = ws1.cell(row=r, column=col, value=val)
            c.alignment = _LEFT if col == 3 else _CENTER
            if fill and row['ng_count'] > 0:
                c.fill = fill

    ws1.column_dimensions['A'].width = 6
    ws1.column_dimensions['B'].width = 10
    ws1.column_dimensions['C'].width = 36
    for col in 'DEFGHIJ':
        ws1.column_dimensions[col].width = 10

    # ── Sheet 2: NG 详细记录 ──
    ws2 = wb.create_sheet('NG详细记录')
    headers2 = ['序号', '报告文件', 'FAI', '描述', '实测值', '名义值', '上公差', '下公差', '偏差']
    for col, h in enumerate(headers2, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = _CENTER

    for i, d in enumerate(stats['details'], 1):
        r = i + 1
        values = [
            i, d['file'], d.get('label', f'FAI_{d["num"]:02d}'), d['desc'],
            d['measured'], d['nominal'], d['upper_tol'], d['lower_tol'], d['deviation'],
        ]
        for col, val in enumerate(values, 1):
            c = ws2.cell(row=r, column=col, value=val)
            c.fill = _DETAIL_FILL
            c.alignment = _LEFT if col in (2, 4) else _CENTER

    ws2.column_dimensions['B'].width = 28
    ws2.column_dimensions['D'].width = 30

    # ── Sheet 3: FAI 明细 ──
    ws3 = wb.create_sheet('FAI明细')
    headers3 = ['FAI', '描述', '名义值', '上公差', '下公差', 'NG次数', 'PASS次数', 'NG率']
    for col, h in enumerate(headers3, 1):
        c = ws3.cell(row=1, column=col, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = _CENTER

    for i, row in enumerate(sorted(stats['summary'], key=lambda x: (x['num'], x.get('sub', ''))), 1):
        r = i + 1
        values = [
            row.get('label', f'FAI_{row["num"]:02d}'), row['desc'],
            row['nominal'], row['upper_tol'], row['lower_tol'],
            row['ng_count'], row['pass_count'], f'{row["ng_rate"]:.1%}',
        ]
        for col, val in enumerate(values, 1):
            ws3.cell(row=r, column=col, value=val).alignment = _CENTER if col != 2 else _LEFT

    wb.save(str(output_path))
    logger.info(f'[NG统计] 已导出: {output_path}')
    return str(output_path)


def format_ng_summary_text(stats: dict) -> str:
    """生成可读的 NG 分析摘要文本。"""
    lines = [
        f'分析报告数: {stats["total_reports"]}',
        f'FAI 项目数: {stats["fai_count"]}',
        f'NG 总次数: {stats["total_ng"]}',
        '',
        'NG 率排行 (前 5):',
    ]
    for i, row in enumerate(stats['summary'][:5], 1):
        if row['ng_count'] == 0:
            continue
        lbl = row.get('label') or f'FAI_{row["num"]:02d}'
        lines.append(
            f'  {i}. {lbl} {row["desc"][:20]} — '
            f'NG {row["ng_count"]}/{row["total"]} ({row["ng_rate"]:.1%}) [{row["severity"]}]'
        )
    if stats['total_ng'] == 0:
        lines.append('  (无 NG 项)')
    return '\n'.join(lines)
