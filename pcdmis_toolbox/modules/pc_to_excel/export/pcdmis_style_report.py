"""PC-DMIS 原生 Excel 报告样式 — 与 FILE / Export / Excel 列布局一致。"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from ..core.models import (
    EXTRA_VALUE_LABELS,
    FeatureRecord,
    PassStatus,
    ReportHeaderInfo,
)
from ..core.tolerance import summarize_results

_HEADER = [
    "尺寸",
    "描述",
    "轴",
    "NOMINAL",
    "MEAS",
    "+TOL",
    "-TOL",
    "BONUS",
    "DEV",
    "OUTTOL",
]

_META_ROWS: tuple[tuple[str, str], ...] = (
    ("PC-DMIS测量程序", "program_name"),
    ("零件名", "part_name"),
    ("序列号", "serial_number"),
    ("日期", "report_date"),
    ("时间", "report_time"),
)

_DATA_HEADER_ROW = len(_META_ROWS) + 2

# 内部轴字段 → 报告轴字母（无 write_axis 时回退）
_AXIS_LETTERS: dict[str, str] = {
    "x": "X",
    "y": "Y",
    "z": "Z",
    "d": "D",
    "length": "L",
    "width": "W",
    "height": "H",
    "angle": "A",
    "minor_d": "MD",
}

_GDT_SYMBOLS = (
    "位置度",
    "POSITION",
    "轮廓度",
    "PROFILE",
    "圆柱度",
    "CYLINDRICITY",
    "圆度",
    "ROUNDNESS",
    "平行度",
    "PARALLELISM",
    "垂直度",
    "PERPENDICULARITY",
    "倾斜度",
    "ANGULARITY",
    "同轴度",
    "CONCENTRICITY",
    "对称度",
    "SYMMETRY",
    "跳动",
    "RUNOUT",
    "全跳动",
    "TOTAL RUNOUT",
)


@dataclass
class PcdmisReportRow:
    dim_id: str
    description: str
    axis: str
    nominal: float | None
    measured: float | None
    plus_tol: float | None
    minus_tol: float | None
    bonus: float | None
    deviation: float | None
    outtol: float | None
    failed: bool = False


def _fmt(value: float | None) -> str | float | int:
    if value is None:
        return ""
    rounded = round(float(value), 4)
    if rounded == int(rounded):
        return int(rounded)
    return rounded


def _split_dim_id(name: str) -> tuple[str, str]:
    """从「ID-特征」拆出尺寸 ID；保留 FAI_2-1 这类带连字符的完整 ID。"""
    text = (name or "").strip()
    if "-" not in text:
        return text, ""
    base, suffix = text.rsplit("-", 1)
    base, suffix = base.strip(), suffix.strip()
    if not base or not suffix:
        return text, ""
    # 仅当后缀含中文（特征名）时剥离，避免 FAI_2-1 → FAI_2
    if any("\u4e00" <= ch <= "\u9fff" for ch in suffix):
        return base, suffix
    return text, ""


def _feature_list(rec: FeatureRecord) -> list[str]:
    feats = [f for f in (rec.feat1, rec.feat2, rec.feat3) if f and f.strip()]
    if feats:
        return feats
    _, suffix = _split_dim_id(rec.name)
    return [suffix] if suffix else []


def _is_gdt_record(rec: FeatureRecord) -> bool:
    ftype = rec.feature_type or ""
    return (
        rec.source_kind in ("tolerance", "fcf")
        or ftype.startswith("TOLERANCE_")
        or (ftype.startswith("FCF_") and ftype != "FCF_SIZE")
    )


def _build_description(rec: FeatureRecord, feats: list[str]) -> str:
    ftype = rec.feature_type or ""
    feat_text = ",".join(feats)

    if _is_gdt_record(rec):
        base = "ISO 几何公差命令"
        return f"{base} ({feat_text})" if feat_text else base

    if ftype in ("FCF_SIZE", "TOLERANCE_SIZE"):
        base = "大小尺寸"
        return f"{base} ({feat_text})" if feat_text else base

    base = ftype if ftype and ftype != "DIMENSION" else "尺寸"
    return f"{base} ({feat_text})" if feat_text else base


def _compute_outtol(
    deviation: float | None,
    plus_tol: float | None,
    minus_tol: float | None,
    stored: float | None,
) -> float | None:
    if stored is not None:
        return stored
    if deviation is None:
        return 0.0 if plus_tol is not None or minus_tol is not None else None
    out = 0.0
    if plus_tol is not None and deviation > plus_tol:
        out = max(out, deviation - plus_tol)
    if minus_tol is not None and deviation < minus_tol:
        out = max(out, minus_tol - deviation)
    return out


def _axis_tol_values(
    rec: FeatureRecord, axis_key: str
) -> tuple[float | None, float | None]:
    plus = rec.plus_tol
    minus = rec.minus_tol
    if rec.tolerance is not None:
        tol_val = getattr(rec.tolerance, axis_key, None)
        if tol_val is not None and plus is None:
            plus = tol_val
        if tol_val is not None and minus is None:
            minus = -abs(tol_val)
    if plus is not None and minus is None:
        minus = -abs(plus)
    return plus, minus


def _axis_letter(rec: FeatureRecord, axis_key: str, multi_axis: bool) -> str:
    if not multi_axis and rec.write_axis:
        return rec.write_axis.strip().upper()
    if axis_key == "d" and "距离" in (rec.feature_type or ""):
        return rec.write_axis.strip().upper() if rec.write_axis else "M"
    return _AXIS_LETTERS.get(axis_key, axis_key.upper())


def _iter_axis_slices(rec: FeatureRecord) -> list[tuple[str, float | None, float | None, float | None]]:
    slices: list[tuple[str, float | None, float | None, float | None]] = []
    extra_keys = tuple(EXTRA_VALUE_LABELS.keys())
    keys = ("x", "y", "z", "d", *extra_keys)

    for key in keys:
        nominal = getattr(rec.nominal, key, None)
        measured = getattr(rec.measured, key, None)
        deviation = getattr(rec.deviation, key, None)
        if nominal is None and measured is None and deviation is None:
            continue
        if deviation is None and nominal is not None and measured is not None:
            deviation = measured - nominal
        slices.append((key, nominal, measured, deviation))

    if not slices and rec.write_axis:
        key = "d"
        for candidate in keys:
            if getattr(rec.measured, candidate, None) is not None:
                key = candidate
                break
        slices.append(
            (
                key,
                getattr(rec.nominal, key, None),
                getattr(rec.measured, key, None),
                getattr(rec.deviation, key, None),
            )
        )
    return slices


def flatten_to_pcdmis_rows(rec: FeatureRecord) -> list[PcdmisReportRow]:
    """将 FeatureRecord 展开为 PC-DMIS 报告行（每轴一行）。"""
    dim_id, _ = _split_dim_id(rec.name)
    feats = _feature_list(rec)
    description = _build_description(rec, feats)
    multi_axis = len(_iter_axis_slices(rec)) > 1
    rows: list[PcdmisReportRow] = []

    for axis_key, nominal, measured, deviation in _iter_axis_slices(rec):
        plus_tol, minus_tol = _axis_tol_values(rec, axis_key)
        outtol = _compute_outtol(deviation, plus_tol, minus_tol, rec.outtol)
        axis = _axis_letter(rec, axis_key, multi_axis)
        if _is_gdt_record(rec) and axis in ("D", "M", ""):
            axis = rec.write_axis.strip().upper() if rec.write_axis else "TP"

        failed = rec.status == PassStatus.FAIL and (
            not rec.failed_axes or axis_key.upper() in rec.failed_axes or axis in rec.failed_axes
        )

        rows.append(
            PcdmisReportRow(
                dim_id=dim_id,
                description=description,
                axis=axis,
                nominal=nominal,
                measured=measured,
                plus_tol=plus_tol,
                minus_tol=minus_tol,
                bonus=rec.bonus,
                deviation=deviation,
                outtol=0.0 if outtol in (None, 0) else outtol,
                failed=failed,
            )
        )

    if not rows:
        plus_tol, minus_tol = _axis_tol_values(rec, "d")
        rows.append(
            PcdmisReportRow(
                dim_id=dim_id,
                description=description,
                axis=rec.write_axis or "D",
                nominal=None,
                measured=None,
                plus_tol=plus_tol,
                minus_tol=minus_tol,
                bonus=rec.bonus,
                deviation=None,
                outtol=rec.outtol,
                failed=rec.status == PassStatus.FAIL,
            )
        )
    return rows


def _row_values(row: PcdmisReportRow) -> list:
    return [
        row.dim_id,
        row.description,
        row.axis,
        _fmt(row.nominal),
        _fmt(row.measured),
        _fmt(row.plus_tol),
        _fmt(row.minus_tol),
        _fmt(row.bonus),
        _fmt(row.deviation),
        _fmt(row.outtol),
    ]


def _write_meta_header(ws, header: ReportHeaderInfo | None) -> None:
    info = header or ReportHeaderInfo()
    for row_idx, (label, field) in enumerate(_META_ROWS, start=1):
        ws.cell(row=row_idx, column=1, value=label)
        ws.cell(row=row_idx, column=2, value=getattr(info, field, "") or "")


def _apply_sheet_style(ws, data_start_row: int) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 32

    for col_idx, title in enumerate(_HEADER, start=1):
        cell = ws.cell(row=_DATA_HEADER_ROW, column=col_idx, value=title)
        cell.font = header_font
        cell.alignment = center
        cell.border = border
        col_letter = get_column_letter(col_idx)
        if col_letter not in {"A", "B"}:
            width = 28 if col_idx == 2 else 10
            ws.column_dimensions[col_letter].width = width

    for row_idx in range(data_start_row, ws.max_row + 1):
        for col_idx in range(1, len(_HEADER) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.alignment = left if col_idx <= 2 else center


def export_pcdmis_excel(
    features: list[FeatureRecord],
    path: Path,
    header: ReportHeaderInfo | None = None,
    part_name: str = "",
    program_path: str = "",
) -> Path:
    if header is None and (part_name or program_path):
        from pathlib import Path as P

        header = ReportHeaderInfo(
            program_name=P(program_path).name if program_path else "",
            part_name=part_name,
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "测量数据"

    _write_meta_header(ws, header)

    for col_idx, title in enumerate(_HEADER, start=1):
        ws.cell(row=_DATA_HEADER_ROW, column=col_idx, value=title)

    data_start_row = _DATA_HEADER_ROW + 1
    row_idx = data_start_row
    for rec in features:
        for row in flatten_to_pcdmis_rows(rec):
            for col_idx, val in enumerate(_row_values(row), start=1):
                ws.cell(row=row_idx, column=col_idx, value=val)
            row_idx += 1

    _apply_sheet_style(ws, data_start_row)

    stats = summarize_results(features)
    footer = ws.max_row + 2
    ws.cell(row=footer, column=1, value="统计")
    ws.cell(row=footer, column=2, value=f"合计 {stats['total']}")
    ws.cell(row=footer, column=4, value=f"合格 {stats['passed']}")
    ws.cell(row=footer, column=6, value=f"超差 {stats['failed']}")
    ws.cell(row=footer, column=8, value=f"合格率 {stats['pass_rate']}%")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_pcdmis_csv(
    features: list[FeatureRecord],
    path: Path,
    header: ReportHeaderInfo | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        info = header or ReportHeaderInfo()
        for label, field in _META_ROWS:
            writer.writerow([label, getattr(info, field, "") or ""])
        writer.writerow([])
        writer.writerow(_HEADER)
        for rec in features:
            for row in flatten_to_pcdmis_rows(rec):
                writer.writerow(_row_values(row))
    return path
