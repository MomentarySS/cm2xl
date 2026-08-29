"""出货检测表按序号填入 — 单元测试。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from ..core.models import AxisValues, FeatureRecord
from ..export.inspection_form_fill import (
    FormFillConfig,
    build_serial_value_maps,
    dim_name_to_serial,
    fill_inspection_form,
    parse_manual_map_text,
    summarize_fill_result,
)


def test_dim_name_to_serial_prefixes():
    assert dim_name_to_serial("FAI_1") == "1"
    assert dim_name_to_serial("CC_1") == "1"
    assert dim_name_to_serial("尺寸_1") == "1"
    assert dim_name_to_serial("FAI_12-1") == "12.1"
    assert dim_name_to_serial("CC_17.2") == "17.2"
    assert dim_name_to_serial("FAI_17A") == "17A"


def test_dim_name_to_serial_manual_map():
    serial = dim_name_to_serial(
        "POS_A",
        prefixes=["FAI_"],
        manual_map={"POS_A": "34"},
    )
    assert serial == "34"


def test_dim_name_custom_prefix_only():
    assert dim_name_to_serial("ABC_9", prefixes=["ABC_"]) == "9"
    assert dim_name_to_serial("FAI_9", prefixes=["ABC_"]) == "FAI_9"


def _make_mini_form(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 1"
    ws["A4"] = "序号"
    ws["G4"] = "测量仪器"
    ws["H4"] = "送检产品序号"
    ws["H5"] = "P1"
    ws["A6"], ws["B6"], ws["G6"], ws["H6"] = 1, 14, "A", 14.0
    ws["A7"], ws["B7"], ws["G7"], ws["H7"] = "12.1", 3, "A", 2.9
    ws["A8"], ws["B8"], ws["G8"], ws["H8"] = 19, 10, "Y", 9.9
    ws["A9"], ws["B9"], ws["G9"], ws["H9"] = 26, 3.2, "CMM", None
    wb.save(path)


def test_fill_only_cmm_rows_by_serial(tmp_path: Path):
    form = tmp_path / "form.xlsx"
    out = tmp_path / "out.xlsx"
    _make_mini_form(form)

    features = [
        FeatureRecord(name="FAI_1", feature_type="尺寸位置", measured=AxisValues(d=14.002)),
        FeatureRecord(name="CC_12-1", feature_type="尺寸位置", measured=AxisValues(d=2.985)),
        FeatureRecord(name="尺寸_19", feature_type="尺寸位置", measured=AxisValues(d=99.0)),
        FeatureRecord(name="FAI_26", feature_type="尺寸位置", measured=AxisValues(d=3.202)),
    ]
    cfg = FormFillConfig(
        id_prefixes=["FAI_", "CC_", "尺寸_"],
        cmm_codes=["A", "CMM"],
        target_col="H",
    )
    result = fill_inspection_form(features, form, out, config=cfg)
    assert result.filled == 3

    wb = load_workbook(out)
    ws = wb["Table 1"]
    assert ws["H6"].value == 14.002
    assert ws["H7"].value == 2.985
    assert ws["H8"].value == 9.9
    assert ws["H9"].value == 3.202


def test_fill_auto_next_empty_column(tmp_path: Path):
    form = tmp_path / "form.xlsx"
    out = tmp_path / "out.xlsx"
    _make_mini_form(form)

    features = [
        FeatureRecord(name="FAI_1", feature_type="尺寸位置", measured=AxisValues(d=14.111)),
        FeatureRecord(name="FAI_12-1", feature_type="尺寸位置", measured=AxisValues(d=2.111)),
    ]
    cfg = FormFillConfig(cmm_codes=["A"], target_col="auto")
    result = fill_inspection_form(features, form, out, config=cfg)
    assert result.target_col == "I"
    wb = load_workbook(out)
    ws = wb["Table 1"]
    assert ws["H6"].value == 14.0
    assert ws["I6"].value == 14.111
    assert ws["I7"].value == 2.111
    assert ws["H8"].value == 9.9


def test_resolve_input_form_chain(tmp_path: Path):
    base = tmp_path / "base.xlsx"
    last = tmp_path / "last.xlsx"
    base.write_bytes(b"x")
    # last 必须是真实可打开的 xlsx：resolve_input_form 现在会校验链式文件完整性
    Workbook().save(last)
    cfg = FormFillConfig(
        form_path=str(base),
        chain_from_last=True,
        last_fill_output=str(last),
    )
    path, chained = cfg.resolve_input_form()
    assert chained and path == last

    cfg.chain_from_last = False
    path, chained = cfg.resolve_input_form()
    assert not chained and path == base


def test_parse_manual_map_text():
    mapping = parse_manual_map_text("FAI_17A=17.1; POS_A:34\n# comment\nCC_1 -> 1")
    assert mapping["FAI_17A"] == "17.1"
    assert mapping["POS_A"] == "34"
    assert mapping["CC_1"] == "1"


def test_fill_piece_id_and_nominal_warning(tmp_path: Path):
    form = tmp_path / "form.xlsx"
    out = tmp_path / "out.xlsx"
    _make_mini_form(form)

    features = [
        FeatureRecord(
            name="FAI_1",
            feature_type="尺寸位置",
            measured=AxisValues(d=14.002),
            nominal=AxisValues(d=99.0),  # 故意与表规格 14 不符
        ),
    ]
    cfg = FormFillConfig(
        cmm_codes=["A"],
        target_col="I",
        write_piece_id=True,
        piece_id="721-009",
        nominal_check=True,
        nominal_tol=0.05,
        manual_map={"FAI_1": "1"},
    )
    result = fill_inspection_form(features, form, out, config=cfg, piece_id="721-009")
    assert result.filled == 1
    assert result.piece_id_written == "721-009"
    assert any("表规格=14" in w for w in result.nominal_warnings)

    wb = load_workbook(out)
    ws = wb["Table 1"]
    assert ws["I5"].value == "721-009"
    assert ws["I6"].value == 14.002


def test_fill_auto_skips_column_with_data_even_if_header_empty(tmp_path: Path):
    """表头空但 CMM 行已有数据时，auto 不得覆盖该列。"""
    form = tmp_path / "form.xlsx"
    out = tmp_path / "out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 1"
    ws["A6"], ws["B6"], ws["G6"], ws["H6"] = 1, 14, "A", 14.0
    ws["A7"], ws["B7"], ws["G7"], ws["H7"] = "12.1", 3, "A", 2.9
    # H5 表头为空
    wb.save(form)

    features = [
        FeatureRecord(name="FAI_1", feature_type="尺寸位置", measured=AxisValues(d=14.111)),
        FeatureRecord(name="FAI_12-1", feature_type="尺寸位置", measured=AxisValues(d=2.111)),
    ]
    result = fill_inspection_form(
        features, form, out, config=FormFillConfig(cmm_codes=["A"], target_col="auto")
    )
    assert result.target_col == "I"
    ws2 = load_workbook(out)["Table 1"]
    assert ws2["H6"].value == 14.0
    assert ws2["I6"].value == 14.111


def test_fill_rejects_same_output_path(tmp_path: Path):
    form = tmp_path / "form.xlsx"
    _make_mini_form(form)
    features = [FeatureRecord(name="FAI_1", feature_type="尺寸位置", measured=AxisValues(d=1.0))]
    try:
        fill_inspection_form(
            features, form, form, config=FormFillConfig(cmm_codes=["A"], target_col="I")
        )
        assert False, "should raise"
    except ValueError as exc:
        assert "拒绝覆盖" in str(exc) or "相同" in str(exc)


def test_preview_form_fill(tmp_path: Path):
    from ..export.inspection_form_fill import preview_form_fill

    form = tmp_path / "form.xlsx"
    _make_mini_form(form)
    features = [
        FeatureRecord(name="FAI_1", feature_type="尺寸位置", measured=AxisValues(d=14.1)),
        FeatureRecord(name="FAI_12-1", feature_type="尺寸位置", measured=AxisValues(d=2.9)),
    ]
    preview = preview_form_fill(
        features,
        form,
        config=FormFillConfig(cmm_codes=["A"], target_col="auto"),
        piece_id="P-9",
        chained=False,
    )
    assert preview.target_col == "I"
    assert preview.will_fill == 2
    assert preview.cmm_rows == 2
    assert "P-9" in preview.summary_text()


def test_multi_axis_prefers_diameter_over_xyz():
    """圆柱 X/Y/D 同时有值时，出货表写 D，不用参考轴 X/Y。"""
    rec = FeatureRecord(
        name="CC_15",
        feature_type="圆柱",
        measured=AxisValues(x=0.0, y=0.0, d=10.192),
        nominal=AxisValues(x=0.0, y=0.0, d=10.0),
    )
    measured, nominal, _ = build_serial_value_maps([rec], FormFillConfig())
    assert measured["15"] == 10.192
    assert nominal["15"] == 10.0


def test_single_axis_angle_still_fills():
    """单轴角度没有 D，仍取 A。"""
    rec = FeatureRecord(
        name="CC_8",
        feature_type="角度",
        measured=AxisValues(angle=89.999),
        nominal=AxisValues(angle=90.0),
        write_axis="A",
    )
    measured, _, _ = build_serial_value_maps([rec], FormFillConfig())
    assert measured["8"] == 89.999


def test_multi_axis_xz_d_prefers_diameter():
    rec = FeatureRecord(
        name="CC_18",
        feature_type="圆柱",
        measured=AxisValues(x=-15.022, z=-7.007, d=10.186),
        nominal=AxisValues(x=-15.0, z=-7.0, d=10.0),
    )
    measured, nominal, _ = build_serial_value_maps([rec], FormFillConfig())
    assert measured["18"] == 10.186
    assert nominal["18"] == 10.0


def test_fill_keeps_serial_conflicts_in_result(tmp_path: Path):
    form = tmp_path / "form.xlsx"
    out = tmp_path / "out.xlsx"
    _make_mini_form(form)
    features = [
        FeatureRecord(name="FAI_1", feature_type="尺寸位置", measured=AxisValues(d=14.0)),
        FeatureRecord(name="CC_1", feature_type="尺寸位置", measured=AxisValues(d=14.5)),
    ]
    result = fill_inspection_form(
        features, form, out, config=FormFillConfig(cmm_codes=["A"], target_col="I")
    )
    assert result.serial_conflicts
    summary = summarize_fill_result(result, extract_count=2)
    assert "序号冲突" in summary
    assert load_workbook(out)["Table 1"]["I6"].value == 14.5
